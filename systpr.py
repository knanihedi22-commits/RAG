import re
import time
import textwrap
from pathlib import Path
import pdfplumber
import chromadb
import openpyxl
import os
from dotenv import load_dotenv
from huggingface_hub import InferenceClient

load_dotenv()  

HF_TOKEN = os.getenv("HF_TOKEN")
LLM_MODEL       = "meta-llama/Llama-3.1-70B-Instruct"
EMBED_MODEL     = "BAAI/bge-small-en-v1.5"   # hosted on HF, no local install
LLM_TEMPERATURE = 0.1
MAX_NEW_TOKENS  = 1024

DATA_ROOT       = Path("DataRoom_FallahTech_PDF")
CHROMA_DIR      = Path("chroma_db")
COLLECTION_NAME = "fallahtech_t1"

CHUNK_SIZE    = 1200
CHUNK_OVERLAP = 240

llm_client   = InferenceClient(model=LLM_MODEL,   token=HF_TOKEN)
embed_client = InferenceClient(model=EMBED_MODEL,  token=HF_TOKEN)


# INGESTION
def extract_pdf(path: Path) -> list[dict]:
    pages = []
    with pdfplumber.open(path) as pdf:
        for i, page in enumerate(pdf.pages, start=1):
            tables = page.extract_tables()
            table_text = ""
            for table in tables:
                for row in table:
                    cleaned = [str(c).strip() if c else "" for c in row]
                    table_text += " | ".join(cleaned) + "\n"

            raw = page.extract_text() or ""
            combined = raw + ("\n[TABLEAU]\n" + table_text if table_text and table_text not in raw else "")

            if combined.strip():
                pages.append({"text": combined.strip(), "source": path.name, "page": str(i)})
    return pages


def extract_xlsx(path: Path) -> list[dict]:
    wb = openpyxl.load_workbook(path, data_only=True)
    sheets = []
    for name in wb.sheetnames:
        rows = []
        for row in wb[name].iter_rows(values_only=True):
            cleaned = [str(c) if c is not None else "" for c in row]
            if any(c.strip() for c in cleaned):
                rows.append(" | ".join(cleaned))
        if rows:
            sheets.append({"text": "\n".join(rows), "source": path.name, "page": name})
    return sheets


def chunk_text(text: str) -> list[str]:
    paragraphs = re.split(r"\n{2,}", text)
    chunks, current = [], ""
    for para in paragraphs:
        para = para.strip()
        if not para:
            continue
        if len(current) + len(para) + 2 <= CHUNK_SIZE:
            current = (current + "\n\n" + para).strip()
        else:
            if current:
                chunks.append(current)
                current = current[-CHUNK_OVERLAP:].strip() + "\n\n" + para
            else:
                for start in range(0, len(para), CHUNK_SIZE - CHUNK_OVERLAP):
                    chunks.append(para[start:start + CHUNK_SIZE])
                current = ""
    if current.strip():
        chunks.append(current.strip())
    return chunks


def ingest(data_root: Path) -> list[dict]:
    all_chunks, chunk_id = [], 0

    for pdf_path in sorted(data_root.rglob("*.pdf")):
        print(f"  PDF : {pdf_path.relative_to(data_root)}")
        for page in extract_pdf(pdf_path):
            for text in chunk_text(page["text"]):
                all_chunks.append({"id": f"c{chunk_id:04d}", "text": text,
                                    "source": page["source"], "page": page["page"]})
                chunk_id += 1

    for xlsx_path in sorted(data_root.rglob("*.xlsx")):
        print(f"  XLSX: {xlsx_path.relative_to(data_root)}")
        for sheet in extract_xlsx(xlsx_path):
            for text in chunk_text(sheet["text"]):
                all_chunks.append({"id": f"c{chunk_id:04d}", "text": text,
                                    "source": sheet["source"], "page": sheet["page"]})
                chunk_id += 1

    print(f"  ✓ {len(all_chunks)} chunks générés")
    return all_chunks


# EMBEDDING & VECTOR STORE


def get_embeddings(texts: list[str]) -> list[list[float]]:
    """Call HF Inference API to embed a batch of texts."""
    return embed_client.feature_extraction(texts)


def build_vector_store(chunks: list[dict]) -> chromadb.Collection:
    CHROMA_DIR.mkdir(exist_ok=True)
    client = chromadb.PersistentClient(path=str(CHROMA_DIR))

    try:
        col = client.get_collection(COLLECTION_NAME)
        if col.count() == len(chunks):
            print(f"  ✓ Collection existante ({col.count()} chunks) — réutilisée")
            return col
        client.delete_collection(COLLECTION_NAME)
    except Exception:
        pass

    col = client.create_collection(COLLECTION_NAME, metadata={"hnsw:space": "cosine"})

    
    batch_size    = 16
    total_batches = (len(chunks) + batch_size - 1) // batch_size
    for i in range(0, len(chunks), batch_size):
        batch      = chunks[i:i + batch_size]
        texts      = [c["text"] for c in batch]
        print(f"  Embedding batch {i // batch_size + 1}/{total_batches} …")
        embeddings = get_embeddings(texts)
        col.add(
            ids        = [c["id"] for c in batch],
            documents  = texts,
            metadatas  = [{"source": c["source"], "page": c["page"]} for c in batch],
            embeddings = embeddings,
        )

    print(f"  ✓ {len(chunks)} chunks indexés")
    return col

# RETRIEVAL

def retrieve(query: str, col: chromadb.Collection, top_k: int = 6) -> list[dict]:
    embedding = get_embeddings([query])[0]
    results   = col.query(query_embeddings=[embedding], n_results=top_k,
                          include=["documents", "metadatas", "distances"])
    return [
        {"text": doc, "source": meta["source"], "page": meta["page"],
         "similarity": round(1 - dist, 4)}
        for doc, meta, dist in zip(
            results["documents"][0], results["metadatas"][0], results["distances"][0]
        )
    ]


# PROMPTS

SYSTEM_PROMPT = """\
Tu es un analyste financier senior spécialisé dans l'instruction de dossiers \
d'investissement pour des fonds de capital-risque.
Tu analyses le dossier de FallahTech SARL, une startup AgriTech tunisienne \
qui demande un financement Série A de 3 000 000 TND.

Ta mission (T1) : évaluer la SOLVABILITÉ et la LIQUIDITÉ en te basant \
EXCLUSIVEMENT sur les extraits fournis.

Règles générales :
1. Cite SOURCE et PAGE pour chaque chiffre utilisé.
2. Si une information est absente du contexte, dis-le explicitement — \
   n'invente jamais de données.
3. Conclus avec un VERDICT structuré : Risque Faible / Moyen / Élevé.
4. Réponds en français, de façon professionnelle.
5. Structure obligatoire : [Analyse] → [Signaux d'alerte] → [Verdict]

Règles de raisonnement numérique — applique-les TOUJOURS avant de conclure :
- Ratio de liquidité courante (actifs courants / passifs courants) :
    > 2      → Liquidité confortable   → Risque Faible
    1 à 2    → Liquidité acceptable    → Risque Moyen
    0,5 à 1  → Tension de trésorerie   → Risque Élevé
    < 0,5    → Insolvabilité imminente → Risque Critique
- Un ratio de 1,56 est SUPÉRIEUR à 1 : l'entreprise couvre ses dettes \
  à court terme. Ne pas le qualifier d'insuffisant.
- Une chute rapide du ratio (ex. 9,43 → 1,56 en un an) est un signal \
  d'alerte sur la TENDANCE, même si le niveau absolu reste acceptable.
- Avant d'écrire ton verdict, vérifie mentalement : \
  "Le chiffre que je cite contredit-il mon affirmation ?" \
  Si oui, corrige ton affirmation.\
"""


def build_user_prompt(question: str, chunks: list[dict]) -> str:
    context = "\n\n---\n\n".join(
        f"[Extrait {i} — Source : {c['source']}, Page : {c['page']}]\n{c['text']}"
        for i, c in enumerate(chunks, 1)
    )
    return (
        f"CONTEXTE DOCUMENTAIRE :\n\n{context}\n\n---\n\n"
        f"QUESTION :\n{question}\n\n"
        "Structure ta réponse : [Analyse] → [Signaux d'alerte] → [Verdict]. "
        "Cite la source et la page entre parenthèses pour chaque chiffre."
    )


# RAG QUERY

def rag_query(question: str, col: chromadb.Collection, top_k: int = 6) -> dict:
    print(f"\n{'='*70}\n  Q : {question}\n{'='*70}")

    chunks = retrieve(question, col, top_k)
    print(f"\n  {len(chunks)} chunks récupérés :")
    for c in chunks:
        print(f"    • [{c['source']} p.{c['page']}] sim={c['similarity']}")

    start    = time.time()
    response = llm_client.chat_completion(
        messages=[
            {"role": "system", "content": SYSTEM_PROMPT},
            {"role": "user",   "content": build_user_prompt(question, chunks)},
        ],
        temperature = LLM_TEMPERATURE,
        max_tokens  = MAX_NEW_TOKENS,
    )
    elapsed = round(time.time() - start, 2)
    answer  = response.choices[0].message.content

    print(f"\n  RÉPONSE ({elapsed}s) :\n")
    for line in answer.split("\n"):
        print(textwrap.fill(line, 70, subsequent_indent="    ") if line else "")

    seen, sources = set(), []
    for c in chunks:
        key = (c["source"], c["page"])
        if key not in seen:
            seen.add(key)
            sources.append({"source": c["source"], "page": c["page"], "similarity": c["similarity"]})

    return {"question": question, "answer": answer, "sources": sources, "elapsed_s": elapsed}


# TEST QUESTIONS

TEST_QUESTIONS = [
    (
        "Analysez l'évolution du ratio de liquidité courante de FallahTech entre 2023 et 2025. "
        "Ce ratio était de 9,43 en 2024 et de 1,56 en 2025. "
        "Quelle est la capacité de remboursement à court terme et quel est le niveau de risque associé ?"
    ),
    (
        "Les dettes fournisseurs ont été multipliées par 15 en un an (300 000 TND en 2025) "
        "alors que la trésorerie est de 510 000 TND. "
        "Effectuez un stress-test : si les fournisseurs exigeaient un remboursement immédiat, "
        "la société pourrait-elle honorer ses engagements ? Quel verdict de risque établissez-vous ?"
    ),
    (
        "En croisant le résultat net 2025 (45 000 TND), l'EBITDA (75 000 TND) et la structure du bilan, "
        "évaluez la solvabilité structurelle sur 2023-2025. "
        "Cette solidité justifie-t-elle une levée Série A de 3 000 000 TND "
        "à une valorisation pré-money de 12 000 000 TND ?"
    ),
]


# MAIN

def main():
    print("\n" + "="*70)
    print("  FallahTech RAG — T1 : Solvabilité & Liquidité  |  Llama-3.1 / HF")
    print("="*70)

    if not DATA_ROOT.exists():
        raise FileNotFoundError(f"Dossier introuvable : {DATA_ROOT}")

    print("\n[1/3] Ingestion …")
    chunks = ingest(DATA_ROOT)

    print("\n[2/3] Embedding & indexation …")
    col = build_vector_store(chunks)

    print("\n[3/3] Questions de test …")
    for q in TEST_QUESTIONS:
        rag_query(q, col)

    print("\n" + "="*70 + "\n  Done.\n" + "="*70)


if __name__ == "__main__":
    main()
